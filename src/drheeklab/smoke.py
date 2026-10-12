"""Frozen-package QA, exclusively in a disposable directory, never user data."""
import tempfile
from pathlib import Path
from PySide6.QtGui import QImage,QColor
from .storage import Storage
from .security import Auth
from .sales import Sales
from .production import Production,CHECKS
from .finance import Finance
from .finance_exports import export_csv,export_excel
from .documents import Documents
from .notifications import Notifications
from .lifecycle import Lifecycle
from .backup import create_backup,verify_backup
from .recovery import RecoverySession
from . import __version__


def run():
    with tempfile.TemporaryDirectory(prefix='drheeklab-qa-') as directory:
        storage=Storage(Path(directory),'test')
        try:
            auth=Auth(storage);auth.create('senha temporária de QA')
            if not auth.verify('senha temporária de QA'):raise RuntimeError('Autenticação não validada')
            image=QImage(32,32,QImage.Format_RGB32);image.fill(QColor('#6d41c7'))
            logo=storage.home/'Marca'/'qa-only.png'
            if not image.save(str(logo)):raise RuntimeError('Codec PNG não disponível')
            storage.set('logo','Marca/qa-only.png')
            sales=Sales(storage);client=sales.save_client({'name':'Cliente temporário de QA'})
            product=sales.products.save({'name':'Peça temporária de QA','category':'functional'})
            variant=sales.products.variants(product)[0]['id']
            production=Production(storage)
            spool=production.catalog.save_filament({'brand':'QA','material':'PLA','color':'Azul','name':'QA','initial_grams':'100','remaining_grams':'100','purchase':'10','reference_per_g':'0.20'})
            source=Path(directory)/'qa.stl';source.write_text('solid qa\nendsolid qa')
            project=production.products.import_media(product,source,'project')
            op=production.create(1,60,[{'variant_id':variant,'project_id':project,'quantity':1,'filaments':[(spool,'10')]}],'QA','0,2 mm / 15%')
            production.start(op)
            production.confirm(op,60,[{'entry_id':production.entries(op)[0]['id'],'consumptions':{spool:'10'}}])
            photos={}
            for view,color in [('front','red'),('back','green'),('detail','blue')]:
                image.fill(QColor(color));photo=Path(directory)/(view+'.png')
                if not image.save(str(photo)):raise RuntimeError('Codec PNG de inspeção indisponível')
                photos[view]=photo
            piece=production.pieces(op)[0]
            production.inspect(op,piece['id'],'approved',{key:{'result':'ok'} for key in CHECKS[piece['category']]},photos,'QA')
            summary=production.complete(op)
            if summary['total_cents']!=210 or summary['good']!=1:raise RuntimeError('Custo/liberação de produção inválidos')
            if production.catalog.get('printers',1)['accumulated_minutes']!=60:raise RuntimeError('Tempo real não consolidado')
            if production.catalog.get('filaments',spool)['remaining_mg']!=90000:raise RuntimeError('Consumo real não consolidado')
            quote=sales.create_quote(client,[{'variant_id':variant,'quantity':1,'unit_price':'30','fulfillment':'stock'}])
            order=sales.approve_quote(quote)
            sales.receive_payment(order,'10','credit_card','0.30')
            sales.receive_payment(order,'20','pix')
            finance=sales.financial(order)
            if finance['balance']!=0 or finance['net_received']!=2970:raise RuntimeError('Valores financeiros não reconciliados')
            docs=sales.documents(order_id=order)
            if len(docs)!=3:raise RuntimeError('Versões de PDF não preservadas')
            for document in docs:
                data=sales.document_path(document['id']).read_bytes()
                if not data.startswith(b'%PDF-') or len(data)<1000:raise RuntimeError('PDF empacotado inválido')
            sales.deliver(order,'deliver','pickup')
            if sales.inventory.summary()[0]['physical']!=0:raise RuntimeError('Baixa física não validada')
            manager=Finance(storage)
            manager.entry('capital_in','100','Aporte de QA','QA')
            bill=manager.create_payable('Frete de QA',7,'5')
            manager.pay(next(i['id'] for i in manager.payables() if i['payable_id']==bill),'5',proof=sales.document_path(docs[0]['id']))
            report=manager.report()
            if report['metrics']['cash_flow_cents']!=12470 or report['metrics']['cost_of_sales_cents']!=210 or report['metrics']['commercial_result_cents']!=2260:raise RuntimeError('Caixa/CPV/resultado não reconciliados')
            if report['reconciliation']['status']!='ok':raise RuntimeError('Reconciliação financeira divergente')
            excel=export_excel(report,Path(directory)/'qa.xlsx');csv=export_csv(report,Path(directory)/'qa.csv')
            from openpyxl import load_workbook
            book=load_workbook(excel,data_only=False)
            try:
                if 'Caixa' not in book.sheetnames or 'Reconciliação' not in book.sheetnames:raise RuntimeError('Excel empacotado inválido')
            finally:book.close()
            if not csv.read_text(encoding='utf-8-sig'):raise RuntimeError('CSV empacotado vazio')
            documents=Documents(storage)
            production.products.import_media(product,Path(directory)/'front.png','photo')
            generated=[documents.emit_finance(),documents.emit_production(op),documents.emit_catalog([{'variant_id':variant,'price':'30'}])]
            from PySide6.QtPdf import QPdfDocument
            for identifier in generated:
                reader=QPdfDocument(None)
                if reader.load(str(documents.path(identifier)))!=QPdfDocument.Error.None_ or reader.pageCount()<1:raise RuntimeError('PDF de relatório/catálogo inválido')
                reader.close()
            original=documents.path(generated[-1]).read_bytes()
            Lifecycle(storage).archive('generated_documents',generated[-1],True,'QA de recuperação')
            Lifecycle(storage).archive('generated_documents',generated[-1],False,'QA de recuperação')
            if documents.path(generated[-1]).read_bytes()!=original:raise RuntimeError('PDF não preservado pela lixeira')
            Notifications(storage).refresh()
            storage.set('theme','light')
            storage.save_draft('home_notes','Rascunho de QA após reinício')
            storage.close()
            storage=Storage(Path(directory),'test')
            if not Auth(storage).verify('senha temporária de QA') or storage.get('theme')!='light' or storage.load_draft('home_notes')!='Rascunho de QA após reinício':
                raise RuntimeError('Preferências/senha/rascunho não preservados após reinício')
            if Finance(storage).report()['metrics']!=report['metrics'] or Documents(storage).path(generated[-1]).read_bytes()!=original:
                raise RuntimeError('Financeiro/documentos não preservados após reinício')
            backup=create_backup(storage)
            if not verify_backup(backup,'test'):raise RuntimeError('Backup não verificado')
            encrypted=create_backup(storage,password='Senha criptográfica de QA')
            if not verify_backup(encrypted,'test',password='Senha criptográfica de QA'):raise RuntimeError('Backup criptografado não verificado')
            transfer=Path(directory)/'another-computer'
            with RecoverySession(encrypted,'test',password='Senha criptográfica de QA') as session:
                if not Auth(session.source).verify('senha temporária de QA'):raise RuntimeError('Senha de backup não verificada')
                session.apply(transfer,confirmed=True)
            transferred=Storage(transfer,'test')
            try:
                if Documents(transferred).path(generated[-1]).read_bytes()!=original:raise RuntimeError('PDF não preservado na transferência')
            finally:transferred.close()
            return {'status':'passed','version':__version__,'checks':['auth','PNG/logo','quote/order','partial payments/fees','versioned PDFs','OP batch/time/consumption','3-photo individual inspection','stock delivery','payables/capital/cash/CPV reconciliation','CSV/Excel export','branded reports/catalog PDFs','notifications/trash','restart persistence','verified backup/transfer','encrypted backup/restore'],'data':'temporary only'}
        finally:
            storage.close()
