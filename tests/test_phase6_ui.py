import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from datetime import date
from PySide6.QtWidgets import QApplication,QDialog,QMessageBox,QFileDialog
from openpyxl import load_workbook
from drheeklab.storage import Storage
from drheeklab.finance import Finance
from drheeklab.phase6_ui import PayableDialog,PayDialog,OwnerDialog,PurchaseDialog,CategoryDialog,MaterialPurchaseDialog
from drheeklab.ui import Window

APP=QApplication.instance() or QApplication([])


def test_finance_forms_partial_payment_owner_funding_and_navigation(tmp_path):
    store=Storage(tmp_path/'data');f=Finance(store)
    form=PayableDialog(f);form.fields['description'].setText('Frete');form.fields['category'].setCurrentIndex(form.fields['category'].findData(7));form.fields['amount'].setText('10.01');form.fields['count'].setText('3');form.generate()
    assert form.installments.listing.rowCount()==3
    form.installments.listing.item(0,0).setText('2');form.installments.listing.item(1,0).setText('3');form.installments.listing.item(2,0).setText('5.01');form.submit()
    assert form.result()==QDialog.Accepted,form.error.text()
    row=f.payables()[0];pay=PayDialog(f,row);pay.fields['amount'].setText('1');pay.fields['fee'].setText('0.10');pay.submit();assert pay.result()==QDialog.Accepted,pay.error.text()
    owner=OwnerDialog(f);owner.fields['amount'].setText('100');owner.fields['description'].setText('Capital');owner.fields['responsible'].setText('Titular');owner.submit();assert owner.result()==QDialog.Accepted,owner.error.text()
    window=Window(store);window.show();window.navigate(4);APP.processEvents()
    page=window.finance_page;assert window.pages.currentWidget()==page and page.tables['payables'].rowCount()==3
    assert 'R$ 98,90' in page.metrics_label.text() and 'Sem divergências' in page.details.toPlainText()
    window.resize(1100,700);APP.processEvents();assert window.grab().save(str(tmp_path/'finance.png'))
    window.close();APP.processEvents()


def test_purchase_and_category_forms_atomic_stock(tmp_path):
    store=Storage(tmp_path/'data');f=Finance(store);form=PurchaseDialog(f)
    for key,value in {'brand':'Marca','color':'Azul','name':'PLA','purchase':'100'}.items():form.fields[key].setText(value)
    form.fields['paid'].setCurrentIndex(form.fields['paid'].findData(True));form.submit();assert form.result()==QDialog.Accepted,form.error.text()
    assert f.report()['metrics']['cash_flow_cents']==-10000
    category=CategoryDialog(f);category.fields['name'].setText('Serviços externos');category.submit();assert category.result()==QDialog.Accepted,category.error.text()
    invalid=PayableDialog(f);invalid.fields['description'].setText('Valor inválido');invalid.fields['amount'].setText('-1');invalid.submit();assert invalid.result()!=QDialog.Accepted
    invalid.reject();store.close()


def test_real_ui_excel_and_csv_exports(tmp_path,monkeypatch):
    store=Storage(tmp_path/'data');f=Finance(store);f.entry('capital_in','10','Aporte','Titular')
    window=Window(store);page=window.finance_page;window.navigate(4)
    monkeypatch.setattr(QMessageBox,'information',lambda *args:QMessageBox.Ok)
    xlsx=tmp_path/'finance.xlsx';monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *args:(str(xlsx),''));page.excel()
    wb=load_workbook(xlsx)
    try:assert 'Caixa' in wb.sheetnames
    finally:wb.close()
    csv=tmp_path/'finance.csv';monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *args:(str(csv),''));page.csv();assert 'Aporte' in csv.read_text(encoding='utf-8-sig')
    window.close();APP.processEvents()
