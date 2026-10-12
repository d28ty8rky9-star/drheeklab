import csv
from datetime import date,timedelta
from decimal import Decimal
import json
import sqlite3
import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtGui import QImage,QColor
from reportlab.pdfgen import canvas
from drheeklab.storage import Storage,MIGRATIONS,SCHEMA_VERSION
from drheeklab.finance import Finance
from drheeklab.sales import Sales
from drheeklab.production import Production,CHECKS
from drheeklab.finance_exports import export_csv,export_excel
from drheeklab.backup import create_backup,verify_backup,restore_backup

APP=QApplication.instance() or QApplication([])
TODAY=date.today().isoformat()

@pytest.fixture
def system(tmp_path):
    storage=Storage(tmp_path/'data');finance=Finance(storage);sales=Sales(storage)
    client=sales.save_client({'name':'Cliente'})
    product=sales.products.save({'name':'Suporte','category':'functional','suggested_price':'30'})
    variant=sales.products.variants(product)[0]['id']
    yield storage,finance,sales,client,product,variant
    try:storage.close()
    except sqlite3.ProgrammingError:pass


def stock_sale(system,qty=2,total_cost='20',price='25',freight='0',paid='0',fee='0'):
    store,finance,sales,client,_,variant=system
    sales.inventory.receive(variant,qty,'Estoque existente','available','Verificado','opening',total_cost=total_cost,reviewer='QA')
    order=sales.create_order(client,[{'variant_id':variant,'quantity':qty,'unit_price':price,'fulfillment':'stock'}],freight)
    if paid!='0':sales.receive_payment(order,paid,'pix',fee)
    sales.deliver(order,'deliver','pickup');return order


def test_cash_sale_cpv_fees_and_capital_are_separate(system):
    store,f,sales,*_=system
    stock_sale(system,paid='20',fee='1')
    f.entry('capital_in','100','Aporte','Titular');f.entry('capital_out','10','Retirada','Titular')
    bill=f.create_payable('Frete pago',7,'5');f.pay(f.payables()[0]['id'],'5')
    r=f.report();m=r['metrics']
    assert m['sold_cents']==5000 and m['received_gross_cents']==2000 and m['received_net_cents']==1900
    assert m['receivable_cents']==3000 and m['cost_of_sales_cents']==2000 and m['expenses_paid_cents']==600
    assert m['commercial_result_cents']==2400 and m['cash_flow_cents']==10400 and m['cash_closing_cents']==10400
    assert m['pieces_sold']==2 and m['ticket_average']=='50.00' and m['commercial_discounts_cents']==1000
    assert m['capital_in_cents']==10000 and m['capital_out_cents']==1000 and m['pending_cost_quantity']==0
    assert r['reconciliation']['status']=='ok'
    assert store.db.execute('SELECT count(*) FROM financial_costs').fetchone()[0]==2
    f.report();assert store.db.execute('SELECT count(*) FROM financial_costs').fetchone()[0]==2


def test_payable_installments_partial_overdue_fees_and_cancel(system):
    _,f,*_=system;yesterday=(date.today()-timedelta(days=1)).isoformat()
    bill=f.create_payable('Fornecedor',9,'10.01','Loja',incurred_on=yesterday,count=3,first_due=yesterday)
    rows=f.payables();assert [r['amount_cents'] for r in sorted(rows,key=lambda r:r['id'])]==[334,334,333]
    first=rows[0];assert first['status']=='overdue';f.pay(first['id'],'1','transfer','0.10')
    assert f.payables()[0]['balance_cents']==234 and f.payables()[0]['status']=='overdue'
    assert f.report()['metrics']['expenses_paid_cents']==110 and f.report()['metrics']['payable_cents']==901
    with pytest.raises(ValueError):f.pay(first['id'],'3')
    with pytest.raises(ValueError):f.cancel_payable(bill,'Correção')
    other=f.create_payable('Sem pagamento',9,'5');f.cancel_payable(other,'Duplicada')
    assert [r for r in f.payables() if r['payable_id']==other][0]['status']=='cancelled'


@pytest.mark.parametrize('bad',[{'amount':'0'},{'amount':'-1'},{'count':121},{'installments':[{'amount':'2'}]},{'incurred_on':'2099-01-01'}])
def test_invalid_obligations_do_not_write(system,bad):
    store,f,*_=system;data={'description':'Conta','category_id':9,'amount':'10',**bad}
    with pytest.raises(ValueError):f.create_payable(**data)
    assert store.db.execute('SELECT count(*) FROM payables').fetchone()[0]==0


def test_category_edits_do_not_reclassify_old_cash_and_result(system):
    _,f,*_=system;bill=f.create_payable('Manutenção',4,'5');f.pay(f.payables()[0]['id'],'5')
    before=f.report()['metrics'];f.save_category('Manutenção renomeada','operating',4)
    assert f.payables()[0]['category']=='Manutenção' and f.report()['metrics']==before
    assert before['expenses_paid_cents']==500 and before['commercial_result_cents']==0
    f.archive_category(4,True,'Não usar')
    with pytest.raises(ValueError):f.create_payable('Nova',4,'1')
    f.archive_category(4,False,'Voltar')
    assert f.create_payable('Nova',4,'1')


def test_owner_reversal_preserves_records_and_excludes_commercial_income(system):
    store,f,*_=system;identifier=f.entry('capital_in','100','Capital','Titular')
    f.reverse_entry(identifier,'Corrigir lançamento','Titular')
    assert f.report()['metrics']['cash_flow_cents']==0 and f.report()['metrics']['sold_cents']==0
    assert store.db.execute('SELECT count(*) FROM finance_entries').fetchone()[0]==2
    with pytest.raises(sqlite3.IntegrityError):f.reverse_entry(identifier,'Outra','Titular')
    assert store.db.execute('SELECT count(*) FROM finance_entries').fetchone()[0]==2
    f.entry('opening_in','50','Saldo inicial comprovado','Titular')
    assert f.report()['metrics']['cash_closing_cents']==5000


def test_purchase_consumption_is_not_a_second_cash_outflow(system):
    store,f,sales,_,product,variant=system
    bill,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    assert f.report()['metrics']['cash_flow_cents']==-10000
    catalog=Production(store).catalog;calc=catalog.save_calculation('Para estoque',1,60,[(spool,'50')],quantity=5);catalog.confirm_actual(calc,60,{spool:'50'})
    r=f.report();assert r['metrics']['cash_flow_cents']==-10000 and r['metrics']['production_cost_cents']==610 and r['metrics']['sold_cents']==0
    assert catalog.get('filaments',spool)['remaining_mg']==950000 and r['metrics']['expenses_paid_cents']==10000
    assert r['metrics']['commercial_result_cents']==0
    with pytest.raises(ValueError):f.cancel_payable(bill,'Cancelar')


def test_integrated_purchase_atomic_and_existing_link_preserves_stock(system,tmp_path):
    store,f,*_=system;data={'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'}
    before=store.db.execute('SELECT count(*) FROM stock_movements').fetchone()[0]
    with pytest.raises(ValueError):f.purchase_filament(data,installments=[{'amount':'99'}])
    assert store.db.execute('SELECT count(*) FROM filaments').fetchone()[0]==0 and store.db.execute('SELECT count(*) FROM stock_movements').fetchone()[0]==before
    spool=Production(store).catalog.save_filament({**data,'remaining_grams':'500','state':'open'});bill=f.create_payable('Compra anterior',1,'100')
    f.link_existing_purchase(bill,spool);assert Production(store).catalog.get('filaments',spool)['remaining_mg']==500000
    with pytest.raises(sqlite3.IntegrityError):f.link_existing_purchase(bill,spool)
    assert Production(store).catalog.get('filaments',spool)['remaining_mg']==500000


def test_proof_pdf_and_images_owned_hashed_backed_up(system,tmp_path):
    store,f,*_=system;path=tmp_path/'proof.pdf';pdf=canvas.Canvas(str(path));pdf.drawString(30,30,'Comprovante QA');pdf.save()
    bill=f.create_payable('Conta',9,'5');payment=f.pay(f.payables()[0]['id'],'5',proof=path)
    proof=store.db.execute('SELECT id FROM finance_proofs WHERE payment_id=?',(payment,)).fetchone()[0]
    owned=f.proof_path(proof);assert owned!=path and owned.is_file();path.unlink()
    backup=create_backup(store);assert verify_backup(backup,'test');owned.write_bytes(b'changed')
    with pytest.raises(ValueError):f.proof_path(proof)
    with pytest.raises(ValueError):create_backup(store)
    root=store.root;restore_backup(store,backup);reopened=Storage(root)
    try:assert Finance(reopened).proof_path(proof).is_file()
    finally:reopened.close()


def test_bad_proof_rolls_back_payment_and_spool(system,tmp_path):
    store,f,*_=system;path=tmp_path/'fake.png';path.write_bytes(b'bad image');bill=f.create_payable('Conta',9,'5')
    with pytest.raises(ValueError):f.pay(f.payables()[0]['id'],'5',proof=path)
    assert store.db.execute('SELECT count(*) FROM payable_payments').fetchone()[0]==0
    data={'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'}
    with pytest.raises(ValueError):f.purchase_filament(data,paid=True,proof=path)
    assert store.db.execute('SELECT count(*) FROM filaments').fetchone()[0]==0 and store.db.execute('SELECT count(*) FROM payables').fetchone()[0]==1


def test_partial_credit_has_no_new_cash_and_cancel_cost_waits_for_return(system):
    store,f,sales,client,_,variant=system;order=stock_sale(system,paid='50')
    sales.cancel(order,'Desistência',future_credit='50')
    r=f.report()['metrics'];assert r['sold_cents']==0 and r['cost_of_sales_cents']==2000 and r['cash_flow_cents']==5000
    next_order=sales.create_order(client,[{'variant_id':variant,'quantity':1,'unit_price':'30','fulfillment':'production'}]);sales.receive_payment(next_order,'30','customer_credit')
    assert f.report()['metrics']['cash_flow_cents']==5000 and f.report()['metrics']['receivable_cents']==0
    pieces=sales.returnable_pieces(order);sales.return_cancelled_pieces(order,[p['id'] for p in pieces],'Recebidas fisicamente')
    assert f.report()['metrics']['cost_of_sales_cents']==0
    assert store.db.execute('SELECT count(*) FROM financial_returns').fetchone()[0]==2
    assert sales.credit_balance(client)==2000


def test_historical_prices_cpv_preserved_after_catalog_edits_and_future_events(system):
    store,f,sales,_,product,variant=system;order=stock_sale(system,paid='20')
    first=f.report()['metrics'];sales.products.save({'name':'Outro nome','category':'decorative','planned_cost':'99','suggested_price':'999'},product)
    assert f.report()['metrics']==first
    assert f.report()['sales'][0]['description']=='Suporte / Padrão'
    future=(date.today()+timedelta(days=5)).isoformat()
    with store.db:store.db.execute('UPDATE payments SET paid_on=?',(future,))
    r=f.report(end=TODAY)['metrics'];assert r['received_gross_cents']==0 and r['receivable_cents']==5000
    with store.db:store.db.execute('UPDATE payments SET paid_on=?',(TODAY,))
    assert f.report()['metrics']==first


def test_period_balances_include_older_obligations_and_opening_cash(system):
    store,f,*_=system;yesterday=(date.today()-timedelta(days=1)).isoformat()
    f.entry('capital_in','100','Capital anterior','Titular',yesterday)
    bill=f.create_payable('Conta antiga',9,'10',incurred_on=yesterday);f.pay(f.payables()[0]['id'],'3')
    r=f.report(TODAY,TODAY)['metrics']
    assert r['cash_opening_cents']==10000 and r['cash_flow_cents']==-300 and r['cash_closing_cents']==9700 and r['payable_cents']==700
    with pytest.raises(ValueError):f.report('2099-01-01',TODAY)


def test_client_product_and_status_filters_keep_consolidated_cash_explicit(system):
    store,f,sales,client,product,variant=system;stock_sale(system,paid='20',fee='1')
    other=sales.save_client({'name':'Outro'});order=sales.create_order(other,[{'description':'Personalizado','quantity':1,'unit_price':'10'}]);sales.receive_payment(order,'10')
    allr=f.report();filtered=f.report(client_id=client)
    assert allr['metrics']['sold_cents']==6000 and filtered['metrics']['sold_cents']==5000
    assert filtered['metrics']['received_gross_cents']==2000 and filtered['metrics']['cash_flow_cents']==2900 and filtered['scoped']
    prod=f.report(product_id=product);assert prod['metrics']['sold_cents']==5000 and prod['metrics']['receivable_cents']==3000
    assert f.report(status='paid')['metrics']['sold_cents']==1000
    assert f.report(status='partial')['metrics']['sold_cents']==5000


def test_material_purchase_increases_stock_once_and_consumption_no_cash(system):
    store,f,sales,*_=system;material=sales.inventory.save_material({'name':'Cola','unit':'ml','balance':'0','reference_per_unit':'0.1'},reason='Cadastro')
    f.save_category('Insumos renomeados','inventory',2)
    bill=f.purchase_material(material,'100','10',paid=True,operation_key='buy-glue')
    assert sales.inventory.get('materials',material)['balance_milli']==100000 and f.report()['metrics']['cash_flow_cents']==-1000
    sales.inventory.consume_material(material,'10','Produção','glue-used')
    assert f.report()['metrics']['cash_flow_cents']==-1000
    with pytest.raises(sqlite3.IntegrityError):f.purchase_material(material,'100','10',paid=True,operation_key='buy-glue')
    assert sales.inventory.get('materials',material)['balance_milli']==90000


def test_exports_reopen_excel_and_csv_literal_text_no_formula(system,tmp_path):
    store,f,*_=system;f.entry('capital_in','1','=HYPERLINK("http://example.org")','Titular')
    report=f.report();csv_path=export_csv(report,tmp_path/'finance.csv');xlsx=export_excel(report,tmp_path/'finance.xlsx')
    with csv_path.open(encoding='utf-8-sig',newline='') as file:data=list(csv.DictReader(file,delimiter=';'))
    assert data[0]['amount_cents']=='100' and data[0]['description'].startswith("'=")
    wb=load_workbook(xlsx,data_only=False)
    try:
        assert 'Caixa' in wb.sheetnames and 'Reconciliação' in wb.sheetnames
        header=[c.value for c in wb['Caixa'][1]];column=header.index('description')+1;cell=wb['Caixa'].cell(2,column)
        assert cell.data_type=='s' and cell.value.startswith('=HYPERLINK')
        assert wb['Caixa'].cell(2,header.index('amount_cents')+1).value==100
        assert all(c.data_type!='f' for ws in wb for row in ws for c in row)
    finally:wb.close()


def test_reconciliation_detects_corruption_instead_of_claiming_ready(system):
    store,f,sales,*_=system;order=stock_sale(system,paid='20')
    assert f.reconcile()['status']=='ok'
    with store.db:store.db.execute('UPDATE payments SET net_cents=1 WHERE order_id=?',(order,))
    assert f.reconcile()['status']=='divergent' and any('líquido' in x for x in f.reconcile()['issues'])


def test_schema5_migration_preserves_sales_and_backup(tmp_path):
    root=tmp_path/'data';path=root/'TESTES'/'Banco de Dados'/'drheeklab.sqlite3';path.parent.mkdir(parents=True)
    with sqlite3.connect(path) as db:
        for version in range(1,6):
            for statement in MIGRATIONS[version]:db.execute(statement)
        db.execute('PRAGMA user_version=5');db.execute("INSERT INTO drafts VALUES('old','preserved','timestamp')")
    store=Storage(root)
    try:
        assert store.db.execute('PRAGMA user_version').fetchone()[0]==SCHEMA_VERSION
        assert store.load_draft('old')=='preserved' and len(Finance(store).categories())==9
        assert list((store.home/'Backups').glob('pre-migration-*.sqlite3')) and store.db.execute('PRAGMA foreign_key_check').fetchall()==[]
    finally:store.close()


def test_actual_op_failure_cpv_and_printer_results_no_double_cash(system,tmp_path):
    store,f,sales,client,product,variant=system
    _,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    p=Production(store);source=tmp_path/'piece.stl';source.write_text('solid test');project=p.products.import_media(product,source,'project')
    order=sales.create_order(client,[{'variant_id':variant,'quantity':1,'unit_price':'30'}]);item=sales.items(order)[0]['id']
    op=p.create(1,60,[{'variant_id':variant,'project_id':project,'quantity':2,'filaments':[(spool,'10')]}],'QA','0,2 mm',item_id=item)
    p.start(op);p.confirm(op,60,[{'entry_id':p.entries(op)[0]['id'],'consumptions':{spool:'10'},'failed':1}],reason='Uma unidade perdida')
    photos={}
    for key,color in [('front','red'),('back','green'),('detail','blue')]:
        image=QImage(30,20,QImage.Format_RGB32);image.fill(QColor(color));path=tmp_path/(key+'.png');assert image.save(str(path));photos[key]=path
    piece=p.pieces(op)[0];p.inspect(op,piece['id'],'approved',{k:{'result':'ok'} for k in CHECKS[piece['category']]},photos,'QA');p.complete(op)
    before=f.report()['metrics'];assert before['production_cost_cents']==210 and before['production_loss_included_cents']==105 and before['cash_flow_cents']==-10000
    assert before['cost_of_sales_cents']==0 and before['pending_cost_quantity']==1
    sales.reserve_item(item);sales.receive_payment(order,'10');sales.deliver(order,'deliver','pickup')
    r=f.report(printer_id=1);m=r['metrics']
    assert m['sold_cents']==3000 and m['cost_of_sales_cents']==210 and m['commercial_result_cents']==2790
    assert m['cash_flow_cents']==-9000 and m['production_cost_cents']==210 and m['pending_cost_quantity']==0
    assert m['received_gross_cents']==1000 and m['receivable_cents']==2000
    assert r['printers'][0]['gross_result_cents']==2790
    assert r['production'][0]['maintenance_reserve_cents']==25 and r['production'][0]['energy_cents']==10 and r['production'][0]['investment_cents']==75
    maintenance=f.create_payable('Manutenção efetiva',4,'20');row=next(i for i in f.payables() if i['payable_id']==maintenance);f.pay(row['id'],'20')
    unfiltered=f.report()['metrics'];assert unfiltered['commercial_result_cents']==2790 and unfiltered['cash_flow_cents']==-11000
    assert f.reconcile()['status']=='ok'


def test_duplicate_payment_key_never_debits_twice(system):
    store,f,*_=system;f.create_payable('Conta',9,'10');row=f.payables()[0]
    f.pay(row['id'],'2',operation_key='pay-once')
    with pytest.raises(sqlite3.IntegrityError):f.pay(row['id'],'2',operation_key='pay-once')
    assert f.payables()[0]['paid_cents']==200 and f.report()['metrics']['cash_flow_cents']==-200


def test_two_printer_receipt_and_product_debt_allocation_conserves_one_cent(system,tmp_path):
    store,f,sales,client,product,variant=system;p=Production(store)
    printer2=p.catalog.save_printer({'manufacturer':'Marca','model':'Outra','identifier':'PR-02'})
    _,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    source=tmp_path/'piece.stl';source.write_text('solid test');project=p.products.import_media(product,source,'project')
    photos={}
    for key,color in [('front','red'),('back','green'),('detail','blue')]:
        image=QImage(30,20,QImage.Format_RGB32);image.fill(QColor(color));path=tmp_path/(key+'.png');assert image.save(str(path));photos[key]=path
    order=sales.create_order(client,[{'variant_id':variant,'quantity':2,'unit_price':'30'}]);item=sales.items(order)[0]['id']
    for printer in (1,printer2):
        op=p.create(printer,60,[{'variant_id':variant,'project_id':project,'quantity':1,'filaments':[(spool,'10')]}],'QA','0,2 mm',item_id=item)
        p.start(op);p.confirm(op,60,[{'entry_id':p.entries(op)[0]['id'],'consumptions':{spool:'10'}}]);piece=p.pieces(op)[0]
        p.inspect(op,piece['id'],'approved',{k:{'result':'ok'} for k in CHECKS[piece['category']]},photos,'QA');p.complete(op)
    sales.reserve_item(item);sales.receive_payment(order,'0.01');sales.deliver(order,'deliver','pickup')
    first=f.report(printer_id=1)['metrics'];second=f.report(printer_id=printer2)['metrics'];product_report=f.report(product_id=product)['metrics']
    assert first['received_gross_cents']+second['received_gross_cents']==1
    assert first['receivable_cents']+second['receivable_cents']==5999
    assert product_report['receivable_cents']==5999 and product_report['received_gross_cents']==1
    assert first['sold_cents']+second['sold_cents']==6000 and f.reconcile()['status']=='ok'


def test_stock_material_loss_is_cost_writeoff_without_new_cash(system):
    store,f,sales,*_=system
    _,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    sales.inventory.discard_material('filament',spool,'10','Material danificado','damaged')
    r=f.report();m=r['metrics']
    assert m['cash_flow_cents']==-10000 and m['expenses_paid_cents']==10000
    assert m['stock_loss_cents']==100 and m['commercial_result_cents']==-100
    assert len(r['stock_losses'])==1 and r['stock_losses'][0]['cost_basis']=='acquisition'


def test_mixed_piece_discard_cost_filters_only_matching_product(system):
    store,f,sales,_,product,variant=system
    _,first=sales.inventory.receive(variant,1,'Inicial','available','Bom','first',total_cost='10',reviewer='QA')
    other_product=sales.products.save({'name':'Outra peça','category':'decorative'});other_variant=sales.products.variants(other_product)[0]['id']
    _,second=sales.inventory.receive(other_variant,1,'Inicial','available','Bom','second',total_cost='20',reviewer='QA')
    sales.inventory.discard_pieces(first+second,'Dano físico','mixed-discard')
    assert f.report()['metrics']['stock_loss_cents']==3000
    assert f.report(product_id=product)['metrics']['stock_loss_cents']==1000
    assert f.report(product_id=other_product)['metrics']['stock_loss_cents']==2000
    assert f.report(product_id=product)['stock_losses'][0]['quantity_milli']==1000


def test_failed_op_before_reprint_is_expense_without_extra_cash_then_absorption_reverses(system,tmp_path):
    store,f,sales,client,product,variant=system;p=Production(store)
    _,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    source=tmp_path/'piece.stl';source.write_text('solid test');project=p.products.import_media(product,source,'project')
    op=p.create(1,60,[{'variant_id':variant,'project_id':project,'quantity':1,'filaments':[(spool,'10')]}],'QA','0,2 mm');p.start(op)
    p.confirm(op,30,[{'entry_id':p.entries(op)[0]['id'],'consumptions':{spool:'5'},'failed':1}],progress='50',reason='Falha completa')
    first=f.report();assert first['metrics']['unabsorbed_production_loss_cents']==105 and first['metrics']['commercial_result_cents']==-105 and first['metrics']['cash_flow_cents']==-10000
    # A distinct quick reprint does not erase the original failed cost.
    new=p.reprint(op,'QA');assert f.report()['metrics']['unabsorbed_production_loss_cents']==105
    # Retrying the original and completing absorbs the old loss in its piece.
    p.start(op);p.confirm(op,60,[{'entry_id':p.entries(op)[0]['id'],'consumptions':{spool:'10'}}])
    photos={}
    for key,color in [('front','red'),('back','green'),('detail','blue')]:
        image=QImage(30,20,QImage.Format_RGB32);image.fill(QColor(color));path=tmp_path/(key+'.png');assert image.save(str(path));photos[key]=path
    piece=p.pieces(op)[0];p.inspect(op,piece['id'],'approved',{k:{'result':'ok'} for k in CHECKS[piece['category']]},photos,'QA');p.complete(op)
    r=f.report();assert r['metrics']['unabsorbed_production_loss_cents']==0 and r['metrics']['commercial_result_cents']==0
    assert [x['amount_cents'] for x in r['production_adjustments']]==[105,-105]
    order=sales.create_order(client,[{'variant_id':variant,'quantity':1,'unit_price':'30','fulfillment':'stock'}]);sales.deliver(order,'deliver','pickup')
    result=f.report()['metrics'];assert result['cost_of_sales_cents']==315 and result['commercial_result_cents']==2685 and result['cash_flow_cents']==-10000


def test_reprint_original_failed_cost_and_new_cpv_both_reconciled(system,tmp_path):
    store,f,sales,client,product,variant=system;p=Production(store)
    _,spool=f.purchase_filament({'brand':'Marca','material':'PLA','color':'Azul','name':'PLA','purchase':'100','initial_grams':'1000','reference_per_g':'0.2'},paid=True)
    source=tmp_path/'piece.stl';source.write_text('solid test');project=p.products.import_media(product,source,'project')
    order=sales.create_order(client,[{'variant_id':variant,'quantity':1,'unit_price':'30'}]);item=sales.items(order)[0]['id']
    original=p.create(1,60,[{'variant_id':variant,'project_id':project,'quantity':1,'filaments':[(spool,'10')]}],'QA','0,2 mm',item_id=item)
    p.start(original);p.confirm(original,30,[{'entry_id':p.entries(original)[0]['id'],'consumptions':{spool:'5'},'failed':1}],progress='50',reason='Falha completa')
    new=p.reprint(original,'QA');p.start(new);p.confirm(new,60,[{'entry_id':p.entries(new)[0]['id'],'consumptions':{spool:'10'}}])
    photos={}
    for key,color in [('front','red'),('back','green'),('detail','blue')]:
        image=QImage(30,20,QImage.Format_RGB32);image.fill(QColor(color));path=tmp_path/(key+'.png');assert image.save(str(path));photos[key]=path
    piece=p.pieces(new)[0];p.inspect(new,piece['id'],'approved',{k:{'result':'ok'} for k in CHECKS[piece['category']]},photos,'QA');p.complete(new)
    sales.reserve_item(item);sales.deliver(order,'deliver','pickup');m=f.report()['metrics']
    assert m['cost_of_sales_cents']==210 and m['unabsorbed_production_loss_cents']==105 and m['commercial_result_cents']==2685
    assert m['production_cost_cents']==315 and m['cash_flow_cents']==-10000
    assert f.report(printer_id=1)['metrics']['commercial_result_cents']==2685
